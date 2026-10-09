"""Import du référentiel maître (Excel) dans l'outil.

Usage :
    python tools/referentiel.py referentiel_BPP_SFPO.xlsx > migrations/0NN_referentiel_<version>.sql

Le fichier maître (feuilles « Mode d'emploi », « Référentiel », « Références BPP », « Journal »)
décrit une ligne par titre ou question, avec niveaux et liens explicites. Le script :
  - contrôle le fichier (codes uniques, chapeaux et conditions existants, références documentées) ;
  - génère une migration SQL qui remplace sections et questions, met à jour les textes des
    références, et fait suivre les réponses déjà saisies grâce à la colonne « Ancien code »
    (renumérotation, fusion). Les réponses à des questions supprimées sont effacées.
La migration refuse de s'appliquer si la base n'est pas dans la version à laquelle se rapporte
la colonne « Ancien code » (cellule « « Ancien code » se rapporte à la version »), pour éviter
d'appliquer deux fois une renumérotation.

Sévérité intrinsèque (critique / majeur / mineur) : colonne « Sévérité » de la feuille
« Référentiel » si elle existe, sinon grille referentiel/grille_severite.csv (code;sévérité;motif).
"""
import csv
import os
import re
import sys

import openpyxl

TYPES = {"Titre", "Question", "Intitulé commun"}
SEVERITES = {"critique", "majeur", "mineur"}
GRILLE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "referentiel", "grille_severite.csv")


def q(s):
    return "NULL" if s is None else "'" + str(s).replace("'", "''") + "'"


def cell_after(ws, label):
    for row in ws.iter_rows(values_only=True):
        if row and row[0] and str(row[0]).strip() == label:
            return str(row[1]).strip() if row[1] is not None else ""
    sys.exit(f"Cellule « {label} » introuvable dans la feuille « {ws.title} »")


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    info = wb["Mode d'emploi"]
    version = cell_after(info, "Version du référentiel")
    previous = cell_after(info, "« Ancien code » se rapporte à la version")

    texts = {}
    for ref, texte, *_ in wb["Références BPP"].iter_rows(min_row=2, values_only=True):
        if ref:
            texts[str(ref).strip()] = (texte or "").strip()

    header = [str(v or "").strip() for v in next(wb["Référentiel"].iter_rows(max_row=1, values_only=True))]
    col_sev = header.index("Sévérité") if "Sévérité" in header else None
    grille = {}
    if col_sev is None and os.path.exists(GRILLE):
        with open(GRILLE, encoding="utf-8") as f:
            grille = {r["code"]: r["severite"] for r in csv.DictReader(f, delimiter=";")}
    rows = [r for r in wb["Référentiel"].iter_rows(min_row=2, values_only=True) if any(v is not None for v in r[:10])]
    errors, sections, questions = [], [], []
    stack = []  # (niveau, id de section)
    codes = {str(r[3]).strip() for r in rows if r[1] != "Titre" and r[3]}
    seen = set()
    for i, r in enumerate(rows, 2):
        _, typ, niveau, code, ancien, chapeau, cond_q, cond_r, intitule, refs = r[:10]
        typ = (typ or "").strip()
        if typ not in TYPES:
            errors.append(f"ligne {i} : type « {typ} » inconnu")
            continue
        intitule = (intitule or "").strip()
        if typ == "Titre":
            try:
                niveau = int(niveau)
            except (TypeError, ValueError):
                errors.append(f"ligne {i} : niveau manquant pour le titre « {intitule[:40]} »")
                continue
            while stack and stack[-1][0] >= niveau:
                stack.pop()
            sid = len(sections) + 1
            sections.append((sid, stack[-1][1] if stack else None, intitule, niveau))
            stack.append((niveau, sid))
            continue
        code = str(code or "").strip()
        if not code:
            errors.append(f"ligne {i} : code manquant")
            continue
        if code in seen:
            errors.append(f"ligne {i} : code {code} en double")
        seen.add(code)
        if not stack:
            errors.append(f"ligne {i} : {code} avant le premier titre")
            continue
        chapeau = str(chapeau).strip() if chapeau else None
        if chapeau and chapeau not in codes:
            errors.append(f"ligne {i} : question chapeau {chapeau} inconnue ({code})")
        cond_q = str(cond_q).strip() if cond_q else None
        cond_r = (str(cond_r).strip().lower() if cond_r else "oui") if cond_q else None
        if cond_q and cond_q not in codes:
            errors.append(f"ligne {i} : condition sur {cond_q} inconnue ({code})")
        if cond_q == code:
            errors.append(f"ligne {i} : {code} dépend d'elle-même")
        if cond_r not in (None, "oui", "non"):
            errors.append(f"ligne {i} : réponse de condition « {cond_r} » (OUI ou NON attendu)")
        ref_list = [x.strip() for x in str(refs or "").split(",") if x.strip()]
        for ref in ref_list:
            if not texts.get(ref):
                errors.append(f"ligne {i} : référence {ref} sans texte dans « Références BPP » ({code})")
        anciens = [x.strip() for x in str(ancien or "").split(",") if x.strip()]
        sev = r[col_sev] if col_sev is not None and col_sev < len(r) else grille.get(code)
        sev = str(sev).strip().lower() if sev else None
        if sev and sev not in SEVERITES:
            errors.append(f"ligne {i} : sévérité « {sev} » (critique, majeur ou mineur attendu) ({code})")
            sev = None
        questions.append(dict(code=code, section=stack[-1][1], chapeau=chapeau, texte=intitule,
                              refs=ref_list, cond_q=cond_q, cond_r=cond_r, part=(typ == "Intitulé commun"),
                              anciens=anciens, sev=sev))
    if errors:
        sys.stderr.write("Référentiel non importé :\n  " + "\n  ".join(errors) + "\n")
        sys.exit(1)

    # correspondance ancien -> nouveau code : explicite (Ancien code), sinon code inchangé
    moved_from = {a for qq in questions for a in qq["anciens"]}
    code_map = [(a, qq["code"]) for qq in questions for a in qq["anciens"]]
    code_map += [(qq["code"], qq["code"]) for qq in questions if qq["code"] not in moved_from]

    out = print
    out(f"-- Référentiel BPP version {version} (généré par tools/referentiel.py depuis {path.split('/')[-1]}).")
    out("-- Remplace sections et questions ; les réponses suivent leur question (colonne « Ancien code »).")
    out("CREATE TABLE IF NOT EXISTS referentiel_info (id INT PRIMARY KEY DEFAULT 1 CHECK (id = 1), "
        "version TEXT NOT NULL, imported_at TIMESTAMPTZ NOT NULL DEFAULT now());")
    out("DO $$ BEGIN")
    out(f"  IF COALESCE((SELECT version FROM referentiel_info), 'initiale') <> {q(previous)} THEN")
    out(f"    RAISE EXCEPTION 'Référentiel en base : %, attendu : {previous.replace(chr(39), chr(39) * 2)}', "
        "COALESCE((SELECT version FROM referentiel_info), 'initiale');")
    out("  END IF;")
    out("END $$;")
    out("")
    out("-- Textes officiels des références")
    out("CREATE TABLE IF NOT EXISTS ref_texts (ref TEXT PRIMARY KEY, texte TEXT NOT NULL);")
    out("DELETE FROM ref_texts;  -- la feuille « Références BPP » fait foi")
    items = sorted(texts.items())
    for k in range(0, len(items), 100):
        out("INSERT INTO ref_texts (ref, texte) VALUES")
        out(",\n".join(f"({q(r)}, {q(t)})" for r, t in items[k:k + 100] if t))
        out("ON CONFLICT (ref) DO UPDATE SET texte = EXCLUDED.texte;")
    out("")
    out("-- Réponses : renommage vers les nouveaux codes (fusion : la réponse la plus récente l'emporte)")
    out("ALTER TABLE responses DROP CONSTRAINT IF EXISTS responses_question_id_fkey;")
    out("CREATE TEMP TABLE code_map (old TEXT PRIMARY KEY, new TEXT NOT NULL) ON COMMIT DROP;")
    for k in range(0, len(code_map), 200):
        out("INSERT INTO code_map (old, new) VALUES")
        out(",\n".join(f"({q(a)}, {q(b)})" for a, b in code_map[k:k + 200]) + ";")
    out("""DELETE FROM responses r USING code_map m
WHERE r.question_id = m.old AND EXISTS (
  SELECT 1 FROM responses r2 JOIN code_map m2 ON m2.old = r2.question_id
  WHERE m2.new = m.new AND r2.evaluation_id = r.evaluation_id AND r2.question_id <> r.question_id
    AND (r2.updated_at, r2.question_id) > (r.updated_at, r.question_id));
UPDATE responses r SET question_id = '~' || m.new FROM code_map m WHERE r.question_id = m.old;
DELETE FROM responses WHERE question_id NOT LIKE '~%';  -- questions supprimées du référentiel
UPDATE responses SET question_id = substr(question_id, 2);""")
    out("")
    out("-- Sections et questions")
    out("ALTER TABLE questions ADD COLUMN IF NOT EXISTS severite TEXT;")
    out("DELETE FROM questions;")
    out("DELETE FROM sections;")
    for k in range(0, len(sections), 100):
        out("INSERT INTO sections (id, parent_id, title, level, sort_order) VALUES")
        out(",\n".join(f"({sid}, {sid_p if sid_p else 'NULL'}, {q(t)}, {lvl}, {sid})"
                       for sid, sid_p, t, lvl in sections[k:k + 100]) + ";")
    out("SELECT setval(pg_get_serial_sequence('sections', 'id'), (SELECT MAX(id) FROM sections));")
    for k in range(0, len(questions), 100):
        out("INSERT INTO questions (id, code, section_id, parent_question_id, question, ref, ref_text, refs, "
            "depends_on_question_id, depends_on_value, is_part, sort_order, severite) VALUES")
        vals = []
        for n, qq in enumerate(questions[k:k + 100], k + 1):
            first = qq["refs"][0] if qq["refs"] else None
            vals.append(f"({q(qq['code'])}, {q(qq['code'])}, {qq['section']}, NULL, {q(qq['texte'])}, {q(first)}, "
                        f"NULL, {q(', '.join(qq['refs']))}, NULL, NULL, "
                        f"{'TRUE' if qq['part'] else 'FALSE'}, {n}, {q(qq['sev'])})")
        out(",\n".join(vals) + ";")
    out("UPDATE questions q SET ref_text = t.texte FROM ref_texts t WHERE t.ref = q.ref;")
    links = [qq for qq in questions if qq["chapeau"] or qq["cond_q"]]
    for qq in links:
        out(f"UPDATE questions SET parent_question_id = {q(qq['chapeau'])}, depends_on_question_id = {q(qq['cond_q'])}, "
            f"depends_on_value = {q(qq['cond_r'])} WHERE id = {q(qq['code'])};")
    out("")
    out("ALTER TABLE responses ADD CONSTRAINT responses_question_id_fkey "
        "FOREIGN KEY (question_id) REFERENCES questions(id) ON DELETE CASCADE;")
    out(f"INSERT INTO referentiel_info (id, version) VALUES (1, {q(version)}) "
        "ON CONFLICT (id) DO UPDATE SET version = EXCLUDED.version, imported_at = now();")
    sys.stderr.write(f"OK : {len(sections)} titres, {len(questions)} questions, {len(texts)} références, "
                     f"{sum(1 for a, b in code_map if a != b)} codes renumérotés.\n")


if __name__ == "__main__":
    main(sys.argv[1])
